from pathlib import Path
import json, re
from pypdf import PdfReader
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]/'rs-solve'
OUT=Path(__file__).resolve().parent
results=[]
terms={'visdrone_paper_tpami.pdf':['standing','person','people'],
 'xview_paper.pdf':['Damaged','hierarch','60 classes'],
 'fhwa_roundabouts.pdf':['Mini-Roundabouts','13 to 25','inscribed'],
 'fiba_basketball_rules_2022.pdf':['15 m','28 m'],
 'ifab_laws_of_the_game_2023_24.pdf':['45 m','68 m','90 m'],
 'dota_paper_tpami.pdf':['gsd','resolution','quadrilateral']}
for p in sorted((ROOT/'reference').glob('*.pdf')):
    item={'file':p.name,'signature':p.read_bytes()[:12].decode('ascii',errors='replace')}
    try:
        doc=PdfReader(p)
        pages=[page.extract_text() or '' for page in doc.pages]
        item.update(pages=len(pages), title=str((doc.metadata or {}).get('/Title','')), text_characters=sum(map(len,pages)))
        (OUT/(p.stem+'.txt')).write_text('\n\n'.join(f'PAGE {i+1}\n'+s for i,s in enumerate(pages)),encoding='utf-8')
        snippets=[]
        for term in terms[p.name]:
            found=0
            for i,s in enumerate(pages):
                pos=s.lower().find(term.lower())
                if pos>=0 and found<3:
                    snippets.append({'page':i+1,'term':term,'text':re.sub(r'\s+',' ',s[max(0,pos-200):pos+650])});found+=1
        item['snippets']=snippets
    except Exception as e: item['error']=type(e).__name__+': '+str(e)
    results.append(item)
    print(json.dumps(item,ensure_ascii=False),flush=True)
p=ROOT/'reference/aircraft_data.xlsx'
wb=load_workbook(p,read_only=True,data_only=True)
excel={'file':p.name,'sheets':[{ 'name':s.title,'rows':s.max_row,'columns':s.max_column} for s in wb]}
sheet=wb['ACD_Data']
rows=list(sheet.values)
excel['header_rows']=[list(r) for r in rows[:2]]
excel['selected_rows']=[]
for i,r in enumerate(rows):
    txt=' '.join(str(x) for x in r if x is not None)
    if any(s in txt.lower() for s in ['cherokee','skyhawk','a320-200','737-800','r22','206b']):
        excel['selected_rows'].append({'excel_row':i+1,'cells':list(r)})
wb.close()
results.append(excel)
(OUT/'reference_checks.json').write_text(json.dumps(results,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(json.dumps(excel,ensure_ascii=False,default=str))
